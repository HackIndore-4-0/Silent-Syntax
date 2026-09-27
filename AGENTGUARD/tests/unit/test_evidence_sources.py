"""New EvidenceSource connectors (Phase 10 breadth):
DocumentEvidenceSource (real PDF/DOCX files), URLEvidenceSource (real
httpx request/response via httpx.MockTransport — no live network),
MongoEvidenceSource (a hand-built fake motor-shaped collection, same
"test our own glue code, not the live database" pattern this suite
already uses for the DeepEval/Ragas adapters), and VectorEvidenceSource
(a caller-supplied async query function, the whole point of that
class's design).
"""
from __future__ import annotations

import pytest

from agentguard.evaluation.dataset import DocumentEvidenceSource, MongoEvidenceSource, URLEvidenceSource, VectorEvidenceSource


class TestDocumentEvidenceSource:
    async def test_retrieves_from_a_real_docx_file(self, tmp_path):
        import docx

        doc = docx.Document()
        doc.add_paragraph("Our refund policy allows returns within 30 days of purchase.")
        doc.add_paragraph("Shipping normally takes 3 to 5 business days.")
        path = tmp_path / "policy.docx"
        doc.save(str(path))

        source = DocumentEvidenceSource([path], chunk_size=200, chunk_overlap=20)
        results = await source.retrieve("what is the refund policy window?")

        assert results
        assert "30 days" in results[0].text
        assert results[0].source_ref == str(path)

    async def test_retrieves_from_a_real_pdf_file(self, tmp_path):
        from fpdf import FPDF

        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.cell(text="Our refund policy allows returns within 30 days of purchase.")
        path = tmp_path / "policy.pdf"
        pdf.output(str(path))

        source = DocumentEvidenceSource([path], chunk_size=200, chunk_overlap=20)
        results = await source.retrieve("refund policy days")

        assert results
        assert "30 days" in results[0].text

    async def test_no_matching_tokens_returns_empty(self, tmp_path):
        path = tmp_path / "notes.txt"
        path.write_text("completely unrelated content about the weather")
        source = DocumentEvidenceSource([path])

        results = await source.retrieve("refund policy")

        assert results == []


class TestURLEvidenceSource:
    async def test_real_httpx_request_response_via_mock_transport(self):
        import httpx

        def handler(request: httpx.Request) -> httpx.Response:
            assert request.url.params["q"] == "refund window?"
            return httpx.Response(200, json={"passages": [{"text": "refunds within 30 days", "score": 0.9}]})

        def extract(data):
            return [(p["text"], p["score"]) for p in data["passages"]]

        source = URLEvidenceSource("https://example.test/search", extract, transport=httpx.MockTransport(handler))
        results = await source.retrieve("refund window?")

        assert len(results) == 1
        assert results[0].text == "refunds within 30 days"
        assert results[0].retrieval_score == 0.9
        assert results[0].source_ref == "https://example.test/search"

    async def test_custom_params_fn_is_used_to_build_the_request(self):
        import httpx

        captured = {}

        def handler(request: httpx.Request) -> httpx.Response:
            captured["query"] = dict(request.url.params)
            return httpx.Response(200, json={"passages": []})

        source = URLEvidenceSource(
            "https://example.test/search",
            lambda data: [],
            params_fn=lambda q: {"search_term": q, "limit": "3"},
            transport=httpx.MockTransport(handler),
        )
        await source.retrieve("hello")

        assert captured["query"] == {"search_term": "hello", "limit": "3"}


class TestMongoEvidenceSource:
    class _FakeCursor:
        def __init__(self, docs):
            self._docs = docs

        def limit(self, n):
            self._docs = self._docs[:n]
            return self

        def __aiter__(self):
            return self._iterate()

        async def _iterate(self):
            for doc in self._docs:
                yield doc

    class _FakeCollection:
        def __init__(self, docs):
            self._docs = docs
            self.last_query = None

        def find(self, query):
            self.last_query = query
            return TestMongoEvidenceSource._FakeCursor(self._docs)

    async def test_retrieve_maps_documents_to_evidence(self):
        fake_collection = self._FakeCollection(
            [{"text": "refunds within 30 days", "score": 0.85}, {"text": "no text field here"}]
        )
        source = MongoEvidenceSource(
            "mongodb://fake", "db", "policies", lambda q: {"$text": {"$search": q}}, score_field="score"
        )
        source._get_collection = lambda: fake_collection  # bypass the real motor client entirely

        results = await source.retrieve("refund window?")

        assert len(results) == 2
        assert results[0].text == "refunds within 30 days"
        assert results[0].retrieval_score == 0.85
        assert results[1].retrieval_score == 1.0  # no score_field value present -> default 1.0, never fabricated higher
        assert fake_collection.last_query == {"$text": {"$search": "refund window?"}}

    async def test_filter_fn_builds_the_real_query_sent(self):
        fake_collection = self._FakeCollection([])
        source = MongoEvidenceSource("mongodb://fake", "db", "policies", lambda q: {"category": q})
        source._get_collection = lambda: fake_collection

        await source.retrieve("refunds")

        assert fake_collection.last_query == {"category": "refunds"}


class TestVectorEvidenceSource:
    async def test_delegates_to_the_caller_supplied_query_fn(self):
        async def query_fn(question):
            return [("chunk about refunds", 0.93), ("unrelated chunk", 0.2)]

        source = VectorEvidenceSource(query_fn, source_label="my_pinecone_index")
        results = await source.retrieve("refund window?")

        assert len(results) == 2
        assert results[0].text == "chunk about refunds"
        assert results[0].retrieval_score == 0.93
        assert results[0].source_ref == "my_pinecone_index"

    async def test_respects_max_results(self):
        async def query_fn(question):
            return [(f"chunk {i}", 1.0 - i * 0.1) for i in range(10)]

        source = VectorEvidenceSource(query_fn, max_results=3)
        results = await source.retrieve("q")

        assert len(results) == 3
